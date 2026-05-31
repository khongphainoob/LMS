"""Multi-format document parser optimized for educational content."""
import os
import re
from dataclasses import dataclass, field
from typing import List, Dict, Any
import frappe


@dataclass
class ParsedDocument:
	"""Output of document parsing."""
	content: str
	sections: List[Dict[str, Any]] = field(default_factory=list)
	metadata: Dict[str, Any] = field(default_factory=dict)
	raw_text: str = ""


class DocumentParser:
	"""Parse various document types into structured text."""

	def parse(self, file_path: str, file_type: str) -> ParsedDocument:
		"""Route to appropriate parser based on file type."""
		if not os.path.exists(file_path):
			raise FileNotFoundError(f"File not found: {file_path}")

		file_type = file_type.lower().strip('.')
		
		try:
			if file_type == 'pdf':
				return self._parse_pdf(file_path)
			elif file_type in ['docx', 'doc']:
				return self._parse_docx(file_path)
			elif file_type in ['pptx', 'ppt']:
				return self._parse_pptx(file_path)
			else:
				return self._parse_text(file_path)
		except Exception as e:
			frappe.log_error(f"Document parsing error for {file_path}: {e}", "RAG Parser")
			return self._parse_text(file_path)

	def _parse_pdf(self, path: str) -> ParsedDocument:
		import pymupdf4llm
		md_text = pymupdf4llm.to_markdown(path)
		sections = self._extract_sections_from_markdown(md_text)
		return ParsedDocument(
			content=md_text,
			sections=sections,
			metadata={"file_type": "pdf"},
			raw_text=md_text
		)

	def _parse_docx(self, path: str) -> ParsedDocument:
		from docx import Document
		doc = Document(path)
		
		content_parts = []
		sections = []
		current_section = {"title": "General", "content": "", "level": 0, "page": 0}
		
		for para in doc.paragraphs:
			text = para.text.strip()
			if not text:
				continue
				
			if para.style.name.startswith('Heading'):
				try:
					level = int(para.style.name.split(' ')[1])
				except Exception:
					level = 1
					
				if current_section["content"]:
					sections.append(current_section)
				
				current_section = {"title": text, "content": text + "\n", "level": level, "page": 0}
				content_parts.append(f"{'#' * level} {text}")
			else:
				current_section["content"] += text + "\n"
				content_parts.append(text)
		
		if current_section["content"]:
			sections.append(current_section)
			
		md_text = "\n\n".join(content_parts)
		return ParsedDocument(
			content=md_text,
			sections=sections,
			metadata={"file_type": "docx"},
			raw_text=md_text
		)

	def _parse_pptx(self, path: str) -> ParsedDocument:
		from pptx import Presentation
		prs = Presentation(path)
		
		content_parts = []
		sections = []
		
		for i, slide in enumerate(prs.slides):
			title = ""
			if slide.shapes.title and slide.shapes.title.text:
				title = slide.shapes.title.text.strip()
			else:
				title = f"Slide {i+1}"
				
			slide_text = []
			for shape in slide.shapes:
				if hasattr(shape, "text") and shape.text:
					if shape != slide.shapes.title:
						slide_text.append(shape.text.strip())
			
			section_content = f"# {title}\n" + "\n".join(slide_text)
			
			# Extract speaker notes
			if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
				notes = slide.notes_slide.notes_text_frame.text.strip()
				if notes:
					section_content += f"\n\nNotes: {notes}"
					
			content_parts.append(section_content)
			sections.append({
				"title": title,
				"content": section_content,
				"level": 1,
				"page": i + 1
			})
			
		md_text = "\n\n---\n\n".join(content_parts)
		return ParsedDocument(
			content=md_text,
			sections=sections,
			metadata={"file_type": "pptx", "pages": len(prs.slides)},
			raw_text=md_text
		)

	def _parse_text(self, path: str) -> ParsedDocument:
		with open(path, 'r', encoding='utf-8', errors='ignore') as f:
			text = f.read()
		sections = self._extract_sections_from_markdown(text)
		return ParsedDocument(
			content=text,
			sections=sections,
			metadata={"file_type": "txt"},
			raw_text=text
		)
		
	def _extract_sections_from_markdown(self, md_text: str) -> List[Dict[str, Any]]:
		"""Basic markdown section extractor based on ATX headings."""
		sections = []
		lines = md_text.split('\n')
		current_section = {"title": "General", "content": "", "level": 0, "page": 0}
		
		for line in lines:
			match = re.match(r'^(#{1,6})\s+(.+)', line)
			if match:
				if current_section["content"].strip():
					sections.append(current_section)
				level = len(match.group(1))
				title = match.group(2).strip()
				current_section = {"title": title, "content": line + "\n", "level": level, "page": 0}
			else:
				current_section["content"] += line + "\n"
				
		if current_section["content"].strip():
			sections.append(current_section)
			
		return sections
